# Builds architecture.pptx (the three stacks and why System.Private.CoreLib is the exception), and a preview of each
# slide as HTML (%TEMP%/architecture-preview-<n>.html: the same shapes, for a look without PowerPoint).
#   python docs/architecture/make-slides.py
# Requires python-pptx (pip install python-pptx).
import html
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = 1240, 697                       # the slide in px (16:9)
EMU = 6858000 / H                      # EMU a px
PT = 72 / 93                           # pt a px (93 px an inch)
FONT = 'Yu Gothic UI'

C = dict(app='F39A3D', ms='3B73B9', msLight='6D9AD3', net='7B5EA7', core='3E2470', host='24508C', os='6B6B6B', hw='8A8A8A',
         foc='C8102E', focLight='E05A47', compat='3AA357', dev='5B9BD5', side='EEF4FB', white='FFFFFF', red='D32F2F', gold='F2C200',
         ink='333333', missing='FDE3E3')


# Where a part's source is public: (chip fill, text color, border, the legend's text).
MARKS = {
    'rs': ('E9F7EC', '1B5E20', '2E9E4A', 'referencesource(MIT)にソースがある'),
    'oss': ('FCE4EC', '880E4F', 'D81B60', 'Microsoft が別に公開(GitHub。Apache 2.0・MIT)'),
    'closed': ('DDDDDD', '666666', None, 'ソースは非公開'),
}


def width(text, size):
    """The text's width in px at size px (full-width characters a size, others about half)."""
    return sum(size if ord(c) > 0x2E7F else size * 0.55 for c in text)


def blend(color, alpha=0.12):
    r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    return '%02X%02X%02X' % tuple(round(255 * (1 - alpha) + v * alpha) for v in (r, g, b))


class Slide:
    def __init__(self, title):
        self.title = title
        self.shapes = []   # (kind, dict)

    def rect(self, x, y, w, h, fill, title='', items=(), color='FFFFFF', size=12, border=None, border_width=0, dash=False,
             shape='rect', align='left', anchor='top', bold=True, column=False, chips_top=None, vertical=False):
        self.shapes.append(('rect', dict(x=x, y=y, w=w, h=h, fill=fill, text=title, color=color, size=size, border=border,
                                         border_width=border_width, dash=dash, shape=shape, align=align, anchor=anchor, bold=bold,
                                         vertical=vertical)))
        # The items as chips under the title, wrapped to the box (a column: one a line).
        top = chips_top if chips_top is not None else (22 if title else 6)
        cx, cy, chip_h, gap, size_i = x + 6, y + top, 17, 3, 10.5
        for item in items:
            item, mark = item if isinstance(item, tuple) else (item, None)
            cw = min(w - 12, width(item, size_i) + 12)
            if column:
                cw = w - 12
            elif cx + cw > x + w - 6:
                cx, cy = x + 6, cy + chip_h + gap
            chip_fill, chip_color, chip_border = (blend(fill) if fill != 'FFFFFF' else 'F5F5F5'), C['ink'], None
            if mark in MARKS:
                chip_fill, chip_color, chip_border = MARKS[mark][:3]
            self.shapes.append(('rect', dict(x=cx, y=cy, w=cw, h=chip_h, fill=chip_fill, text=item,
                                             color=chip_color, size=size_i, border=chip_border, border_width=2 if chip_border else 0,
                                             dash=False, shape='round', align='left', anchor='middle', bold=True)))
            if column:
                cy += chip_h + gap
            else:
                cx += cw + gap

    def text(self, x, y, w, h, text, size=13, color=C['ink'], bold=False, align='left'):
        self.shapes.append(('rect', dict(x=x, y=y, w=w, h=h, fill=None, text=text, color=color, size=size, border=None, border_width=0,
                                         dash=False, shape='rect', align=align, anchor='top', bold=bold)))

    def badge(self, x, y, w, h, text, fill=C['red'], color='FFFFFF', size=11):
        self.rect(x, y, w, h, fill, text, color=color, size=size, shape='round', align='center', anchor='middle', border='FFFFFF', border_width=1.5)

    def arrow(self, x1, y1, x2, y2, color=C['ink'], width_px=2.5, dash=False):
        self.shapes.append(('arrow', dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width_px, dash=dash)))


# ---- the PowerPoint ----

def set_font(run, size, color, bold):
    run.font.size = Pt(size * PT)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', FONT)


def to_pptx(slides, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(W * EMU)), Emu(int(H * EMU))
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        all_shapes = [('rect', dict(x=12, y=4, w=W - 24, h=32, fill=None, text=s.title, color=C['ink'], size=21, border=None, border_width=0,
                                    dash=False, shape='rect', align='left', anchor='middle', bold=True))] + s.shapes
        for kind, d in all_shapes:
            if kind == 'arrow':
                line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(d['x1'] * EMU)), Emu(int(d['y1'] * EMU)),
                                                  Emu(int(d['x2'] * EMU)), Emu(int(d['y2'] * EMU)))
                line.line.color.rgb = RGBColor.from_string(d['color'])
                line.line.width = Pt(d['width'] * PT)
                ln = line.line._get_or_add_ln()
                ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
                if d['dash']:
                    line.line.dash_style = 4
                continue
            kind_shape = MSO_SHAPE.ROUNDED_RECTANGLE if d['shape'] == 'round' else MSO_SHAPE.RECTANGLE
            shape = slide.shapes.add_shape(kind_shape, Emu(int(d['x'] * EMU)), Emu(int(d['y'] * EMU)), Emu(int(d['w'] * EMU)), Emu(int(d['h'] * EMU)))
            shape.shadow.inherit = False
            if d['shape'] == 'round':
                shape.adjustments[0] = 0.18
            if d['fill']:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(d['fill'])
            else:
                shape.fill.background()
            if d['border']:
                shape.line.color.rgb = RGBColor.from_string(d['border'])
                shape.line.width = Pt(d['border_width'] * PT)
                if d['dash']:
                    shape.line.dash_style = 7
            else:
                shape.line.fill.background()
            tf = shape.text_frame
            tf.word_wrap = not (d['shape'] == 'round' and d['h'] <= 24)   # a chip or a badge: one line
            tf.auto_size = None
            for side in ('margin_left', 'margin_right'):
                setattr(tf, side, Emu(int(5 * EMU)))
            tf.margin_top = tf.margin_bottom = Emu(int(2 * EMU))
            tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE}[d['anchor']]
            if d.get('vertical'):
                tf._txBody.bodyPr.set('vert', 'eaVert')
            for i, line in enumerate(d['text'].split('\n')):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER}[d['align']]
                run = p.add_run()
                run.text = line
                set_font(run, d['size'], d['color'], d['bold'])
    prs.save(path)


# ---- the preview ----

def to_html(slide, path):
    out = [f'<!DOCTYPE html><html lang="ja"><head><meta charset="utf-8"><style>body{{margin:0;font-family:"{FONT}","Meiryo",sans-serif}}'
           f'#s{{position:relative;width:{W}px;height:{H}px;overflow:hidden}}div.b{{position:absolute;box-sizing:border-box;padding:2px 5px;overflow:hidden;'
           f'white-space:pre-wrap;line-height:1.2}}</style></head><body><div id="s">',
           f'<div class="b" style="left:12px;top:4px;width:{W - 24}px;height:32px;font-size:21px;font-weight:bold;color:#{C["ink"]};display:flex;align-items:center">{html.escape(slide.title)}</div>']
    svg = []
    for kind, d in slide.shapes:
        if kind == 'arrow':
            svg.append(f'<line x1="{d["x1"]}" y1="{d["y1"]}" x2="{d["x2"]}" y2="{d["y2"]}" stroke="#{d["color"]}" stroke-width="{d["width"]}" '
                       f'{"stroke-dasharray=\"6 4\"" if d["dash"] else ""} marker-end="url(#m{d["color"]})"/>')
            continue
        style = [f'left:{d["x"]}px', f'top:{d["y"]}px', f'width:{d["w"]}px', f'height:{d["h"]}px', f'font-size:{d["size"]}px',
                 f'color:#{d["color"]}', f'font-weight:{"bold" if d["bold"] else "normal"}', f'text-align:{d["align"]}']
        if d['fill']:
            style.append(f'background:#{d["fill"]}')
        if d['border']:
            style.append(f'border:{d["border_width"]}px {"dashed" if d["dash"] else "solid"} #{d["border"]}')
        if d['shape'] == 'round':
            style.append('border-radius:4px')
        if d['anchor'] == 'middle':
            style.append('display:flex;align-items:center;' + ('justify-content:center' if d['align'] == 'center' else ''))
        if d.get('vertical'):
            style.append('writing-mode:vertical-rl')
        out.append(f'<div class="b" style="{";".join(style)}">{html.escape(d["text"])}</div>')
    colors = {d['color'] for k, d in slide.shapes if k == 'arrow'}
    markers = ''.join(f'<marker id="m{c}" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="#{c}"/></marker>' for c in colors)
    out.append(f'<svg width="{W}" height="{H}" style="position:absolute;left:0;top:0"><defs>{markers}</defs>{"".join(svg)}</svg></div></body></html>')
    with open(path, 'w', encoding='utf-8') as f:
        f.write(''.join(out))


# ---- the slides ----

def legend(s, entries):
    x = 14
    for color, label in entries:
        s.rect(x, 668, 14, 14, color)
        s.text(x + 16, 664, width(label, 12) + 16, 22, label, size=12)
        x += width(label, 12) + 44


MAIN_W, SIDE_X, SIDE_W = 1048, 1058, 176
# The vertical bands at the left: what belongs to the parts of several layers (language packs, translations, the culture
# profile), and, beside them, what carries the process (IIS, systemd, Docker: it loads the runtime, it is not called by it).
BAND_X, BAND_W, HOST_X, HOST_W_BAND, LEFT = 6, 30, 40, 30, 74
# The rule of the stacks: a box touches the one under it when it uses it directly, and does not when it does not (what
# everything uses, the class libraries and CoreLib, may be used across a layer).
# The application's kinds, a column each: what is under a column is what that kind uses; a box across the columns is
# shared by them all.
COLS = [(LEFT, 190), (268, 108), (380, 104), (488, 96), (588, 108)]
# The application's parts all its kinds share, at the right, each step on what it uses directly: System.Web's core (step 0,
# from STEP0), the libraries only the application uses (APPLIBS), the libraries it shares with System.Web (from STEP1 to
# SHARED_END), then Compat or the class libraries (from STEP2).
STEP0 = 700
CORE_END = 748
APPLIBS_X, APPLIBS_W = 752, 122
STEP1 = 878
SHARED_END = 908
STEP2 = 912
HOST_W = 446                              # ③: the host (beside the libraries, under System.Web's core)
NOTE = ('見方: 上下に接する箱は、上が下を直接使う(接しなければ使わない。どれもが使うクラスライブラリと CoreLib へは、層を飛ばしても使う)。'
        '縦に並ぶ箱は、その上のアプリの種類が使う部品。右のアプリは、段ごとに使う層に載る。左の縦の帯は、いくつもの層に付くもの'
        '(プロセスを載せるホストも)。黄色の印は要求を受ける HTTP サーバー。')


def rows(*spec):
    """The rows from the top: (name, height) -> {name: (y, h)}."""
    out, y = {}, 40
    for name, h in spec:
        out[name] = (y, h)
        y += h + 4
    return out


def bottom(r):
    return r[0] + r[1]


def kinds(s, r, boxes, cols=COLS):
    """A box a column (color, title, items, options)."""
    y, h = r['KINDS']
    for (x, w), (color, title, items, *o) in zip(cols, boxes):
        o = o[0] if o else {}
        if o.get('missing'):
            s.rect(x, y, w, h, 'FFFFFF', title, items, color=C['red'], border=C['red'], border_width=2, dash=True, column=True)
        else:
            s.rect(x, y, w, h, color, title, items, column=o.get('column', True))


def app_kinds(s, r, boxes, cols=COLS):
    y, h = r['APP']
    for (x, w), (color, title, items) in zip(cols, boxes):
        s.rect(x, y, w, h, color, title, items)


def application(s, r, title, items, compat=None):
    """The application's shared parts, at the right: on System.Web's core and the libraries only it uses (step 0), on
    the libraries it shares with System.Web (step 1), then on the class libraries directly, or (compat: its title and
    items) on Compat for the calls whose result differs, beside a strip down to the class libraries for the others."""
    top, kinds_end, core_end = r['APP'][0], bottom(r['KINDS']), bottom(r['CORE'])
    s.rect(STEP0, top, MAIN_W - STEP0, kinds_end - top, C['app'], title, items, column=True, chips_top=38)
    if 'LIBS' not in r:
        s.rect(STEP1, kinds_end, MAIN_W - STEP1, core_end - kinds_end, C['app'])
        s.text(STEP1, core_end - 20, MAIN_W - STEP1, 20, '直接 ↓', size=11, color='FFFFFF', bold=True)
        return
    libs_end = bottom(r['LIBS'])
    s.rect(STEP1, kinds_end, MAIN_W - STEP1, core_end + 4 - kinds_end, C['app'])
    s.text(STEP1, core_end - 20, MAIN_W - STEP1, 20, '下の層も直接使う ↓', size=11, color='FFFFFF', bold=True)
    if not compat:
        s.rect(STEP2, core_end + 4, MAIN_W - STEP2, libs_end - core_end - 4, C['app'])
        s.text(STEP2, libs_end - 20, MAIN_W - STEP2, 20, '直接 ↓', size=11, color='FFFFFF', bold=True)
        return
    compat_h = 100
    s.rect(STEP2, core_end + 4, MAIN_W - STEP2, libs_end - compat_h - 4 - core_end - 4, C['app'])
    s.text(STEP2, libs_end - compat_h - 24, MAIN_W - STEP2, 20, 'Compat ↓(違う所だけ)', size=11, color='FFFFFF', bold=True)
    s.rect(MAIN_W - 28, libs_end - compat_h - 4, 28, compat_h + 4, C['app'], '直接', color='FFFFFF', size=10, align='center',
           anchor='middle', vertical=True)
    s.rect(STEP2, libs_end - compat_h, MAIN_W - 32 - STEP2, compat_h, C['compat'], compat[0], compat[1], column=True, chips_top=36, size=11)


def core(s, r, color, title, items):
    s.rect(LEFT, r['CORE'][0], CORE_END - LEFT, r['CORE'][1], color, title, items)


def app_libs(s, r, color, title, items):
    """The libraries only the application uses (not System.Web): under the application, beside System.Web's core."""
    s.rect(APPLIBS_X, r['CORE'][0], APPLIBS_W, r['CORE'][1], color, title, items, column=True)


def shared_libs(s, r, color, title, items, x0=LEFT):
    """The libraries System.Web and the application both use: under System.Web's core and the application's step 1."""
    s.rect(x0, r['LIBS'][0], SHARED_END - x0, r['LIBS'][1], color, title, items)


def full(s, row, color, title, items=(), x0=LEFT, **options):
    s.rect(x0, row[0], MAIN_W - x0, row[1], color, title, items, **options)


def halves(s, row, left, right, x0=LEFT, split=(LEFT + MAIN_W) // 2):
    """Two boxes in a row: (color, title, items) each."""
    s.rect(x0, row[0], split - x0, row[1], *left)
    s.rect(split + 4, row[0], MAIN_W - split - 4, row[1], *right)


def corelib(s, r, title, items, exception=True):
    """CoreLib with the runtime: one block, under the class libraries (the exception: not replaceable)."""
    y, h = r['CORELIB']
    s.rect(LEFT, y, MAIN_W - LEFT, h, C['core'], title, items, border=C['gold'] if exception else None, border_width=4 if exception else 0)
    if exception:
        s.badge(MAIN_W - 156, y + 5, 146, 20, '例外: 差し替え不可', size=11)


def http_server(s, x, y, text='HTTP サーバー(要求を受ける)'):
    """The mark of what receives the requests (IIS, Kestrel): the same on every slide, wherever the layer it is in."""
    s.badge(x, y, width(text, 11) + 16, 20, text, fill=C['gold'], color=C['ink'], size=11)


def band(s, top, end, title, color, missing=False):
    """A vertical band along the layers whose parts it belongs to."""
    if missing:
        s.rect(BAND_X, top, BAND_W, end - top, 'FFFFFF', title, color=C['red'], border=C['red'], border_width=2, dash=True, size=12,
               align='center', anchor='middle', vertical=True)
    else:
        s.rect(BAND_X, top, BAND_W, end - top, color, title, size=12, align='center', anchor='middle', vertical=True)


def host_band(s, r, title, http_server=False):
    """What carries the process (loads the runtime and the application, gives them the requests): a band along the
    layers in the process, beside them, not under CoreLib (CoreLib does not call it; the runtime uses the OS directly).
    http_server: it receives the requests itself (IIS on ①), the gold border of the mark."""
    top, end = r['CORE'][0], bottom(r['CORELIB'])
    s.rect(HOST_X, top, HOST_W_BAND, end - top, C['host'], title, size=11, align='center', anchor='middle', vertical=True,
           border=C['gold'] if http_server else None, border_width=4 if http_server else 0)


def side(s, color, title, items, marks=False):
    s.rect(SIDE_X, 40, SIDE_W, 604, C['side'])
    s.rect(SIDE_X + 6, 46, SIDE_W - 12, 200 if marks else 400, color, title, items, column=True)
    if marks:
        # The legend of the marks: where each part's source is public.
        s.rect(SIDE_X + 6, 252, SIDE_W - 12, 206, 'FFFFFF', 'ソースの公開(項目の色と枠)', color=C['ink'], border='BBBBBB', border_width=1)
        for i, (key, (fill, ink, border, label)) in enumerate(MARKS.items()):
            y = 278 + i * 60
            s.rect(SIDE_X + 12, y, 44, 18, fill, '項目', color=ink, size=10.5, border=border, border_width=2 if border else 0, shape='round',
                   anchor='middle')
            s.text(SIDE_X + 60, y - 3, SIDE_W - 68, 58, label, size=10.5)
    s.text(SIDE_X + 4, 462 if marks else 452, SIDE_W - 8, 190, NOTE, size=10.5)


APP_KINDS = [
    (C['app'], 'Web Forms', ['.aspx・.ascx・.master', 'コードビハインド']),
    (C['app'], 'MVC 5', ['Controller', '.cshtml(Razor)']),
    (C['app'], 'Web API 2', ['ApiController', 'JSON・XML']),
    (C['app'], 'ASMX', ['.asmx(SOAP)']),
    (C['app'], 'WCF サービス', ['.svc', 'serviceModel']),
]
CORE_ITEMS = ['HttpApplication・パイプライン', 'HTTP モジュール・ハンドラー', 'ルーティング(System.Web.Routing)', 'セッション・キャッシュ',
              'フォーム認証・メンバーシップ・ロール', 'web.config のセクション']


def slide1():
    s = Slide('① Windows Server + .NET Framework 4.8(元のサーバー)')
    r = rows(('APP', 64), ('KINDS', 104), ('CORE', 84), ('LIBS', 44), ('CLASSLIBS', 44), ('CORELIB', 50), ('OS', 56), ('HW', 30))
    app_kinds(s, r, APP_KINDS)
    kinds(s, r, [
        (C['ms'], 'Web Forms の部品', [('System.Web.UI(ページ)', 'rs'), ('ASP.NET AJAX(Extensions)', 'rs'),
                                         ('Chart・Dynamic Data・Mobile', 'rs'), ('バンドル(Optimization)', 'oss')]),
        (C['msLight'], 'MVC 5(NuGet)', [('System.Web.Mvc', 'oss'), ('Razor・WebPages', 'oss'), ('AspNetWebStack', 'oss'), ('バンドル', 'oss')]),
        (C['msLight'], 'Web API 2', [('Web.Http', 'oss'), ('Http.WebHost', 'oss'), 'Newtonsoft.Json']),
        (C['ms'], 'ASMX', [('Web.Services', 'rs'), ('WSDL・SOAP', 'rs')]),
        (C['ms'], 'WCF', [('ServiceModel', 'rs'), ('Activation(.svc)', 'rs'), ('IIS でのホスト', 'rs')]),
    ])
    application(s, r, 'お客様のアプリ\n全種類で共通', ['Global.asax', 'web.config', 'App_Code', 'クラスライブラリ(C#・VB)', '参照する DLL・NuGet'])
    core(s, r, C['ms'], 'System.Web のコア ― 全種類で共通(Microsoft)', [(i, 'rs') for i in CORE_ITEMS + ['実行時コンパイル(Global.asax・App_Code・ページ)']])
    app_libs(s, r, C['ms'], 'アプリのデータ', [('EF6(NuGet)', 'oss'), ('LINQ to SQL', 'rs')])
    shared_libs(s, r, C['ms'], 'ライブラリ ― System.Web もアプリも使う', [('構成(System.Configuration)', 'rs'), ('System.Drawing(.NET 版で公開)', 'oss')])
    full(s, r['CLASSLIBS'], C['net'], '.NET Framework 4.8 クラスライブラリ', [('System', 'rs'), ('System.Data(ADO.NET・SqlClient)', 'rs'),
                                                                         ('System.Xml(XmlSerializer)', 'rs'), ('System.Core', 'rs'),
                                                                         ('System.Runtime.Serialization', 'rs')])
    corelib(s, r, 'mscorlib + CLR 4(一組で動く)', [('mscorlib: object・string・double・CultureInfo・BinaryFormatter', 'rs'),
                                                ('CLR 4(clr.dll): JIT・GC・double の書式', 'closed')], exception=False)
    full(s, r['OS'], C['os'], 'Windows Server', [(i, 'closed') for i in ['IIS の機能(アプリケーション プール・URL Rewrite・静的ファイル・customHeaders・W3C ログ・Windows 認証)',
                                                                          'http.sys', 'NTFS', 'NLS(カルチャのデータ)', 'タイムゾーン', 'GDI+', 'SChannel']])
    full(s, r['HW'], C['hw'], 'サーバー(x64 の物理マシン・仮想マシン)', anchor='middle')
    band(s, r['KINDS'][0], bottom(r['CORELIB']), '言語パック(サテライト。非公開)', C['msLight'])
    host_band(s, r, 'IIS(ネイティブ)が .NET を載せる(w3wp)\n統合パイプラインで System.Web を呼ぶ', http_server=True)
    side(s, C['dev'], '開発・配置', [('Visual Studio', 'closed'), ('MSBuild', 'oss'), ('Web Deploy', 'closed'), ('IIS マネージャー', 'closed'),
                                    ('SQL Server', 'closed')], marks=True)
    legend(s, [(C['app'], 'お客様のアプリ'), (C['ms'], 'Microsoft(.NET Framework)'), (C['msLight'], 'Microsoft(NuGet・言語パック)'), (C['net'], '.NET Framework の実行環境'),
               (C['host'], 'Web サーバー(IIS)'), (C['os'], 'OS')])
    return s


def slide2():
    s = Slide('② Windows / Linux サーバー + .NET 10(そのまま移すと)')
    r = rows(('APP', 64), ('KINDS', 110), ('CORE', 104), ('CLASSLIBS', 44), ('CORELIB', 50), ('OS', 66), ('HW', 30))
    app_kinds(s, r, [
        (C['app'], 'Web Forms', ['動かない', '作り直し(Blazor など)']),
        (C['app'], 'MVC 5', ['Core MVC に', '書き直し']),
        (C['app'], 'Web API 2', ['Core に', '書き直し']),
        (C['app'], 'ASMX', ['動かない']),
        (C['app'], 'WCF サービス', ['CoreWCF に移す']),
    ])
    kinds(s, r, [
        (None, 'System.Web.UI は無い', ['AJAX・Chart・Dynamic Data・', 'Mobile・バンドルも無い'], {'missing': True}),
        (C['ms'], 'Core MVC', ['Razor(新しい版)', 'バンドルは別']),
        (C['ms'], 'ASP.NET Core', ['Minimal API など', 'Text.Json']),
        (None, 'ASMX は無い', [], {'missing': True}),
        (C['msLight'], 'CoreWCF', ['.svc・basicHttp', '一部の機能は無い']),
    ])
    application(s, r, 'お客様のアプリ\n全種類で共通(作り直し)', ['Global.asax → Program.cs', 'web.config → JSON 構成', 'App_Code は動かない', 'ライブラリを .NET 向けに'])
    core(s, r, C['ms'], 'ASP.NET Core ― 全種類の土台(System.Web とは別物。System.Web のコアは無い)', ['Kestrel(.NET の上のライブラリ)', 'ミドルウェア', 'ルーティング(新)',
                                                                                     'セッション(仕組みが違う)', 'Cookie 認証・Identity', '構成(appsettings.json)',
                                                                                     '実行時コンパイル無し'])
    http_server(s, CORE_END - 190, r['CORE'][0] + 4)
    app_libs(s, r, C['ms'], '.NET のライブラリ', ['SqlClient', 'EF Core・EF6', 'Drawing.Common', 'LINQ to SQL 無し'])
    full(s, r['CLASSLIBS'], C['net'], '.NET 10 クラスライブラリ', ['System.Data(DataTable)', 'System.Xml(XmlSerializer)', 'System.Text.Json', 'System.Net.Http',
                                                                'BinaryFormatter(削除・例外)'])
    corelib(s, r, 'System.Private.CoreLib + CoreCLR(一組)', ['CoreLib: object・string・double・CultureInfo', 'double の書式・Parse・カルチャは .NET の動き'])
    halves(s, r['OS'], (C['os'], 'Windows Server', ['IIS + ANCM(Kestrel に転送・.NET を載せる)', 'NTFS', 'NLS または ICU', 'タイムゾーン', 'GDI+']),
           (C['os'], 'Linux', ['systemd・Docker', 'nginx・Apache(任意。Kestrel に転送)', 'ext4・xfs(大文字小文字を区別)', 'ICU', 'IANA のタイムゾーン', 'GDI+ なし']))
    full(s, r['HW'], C['hw'], 'サーバー・コンテナ・クラウド(x64・ARM64)', anchor='middle')
    band(s, r['KINDS'][0], bottom(r['CORELIB']), '言語パックは無い(英語)', None, missing=True)
    host_band(s, r, 'プロセスを載せる\nIIS・systemd・Docker')
    side(s, C['dev'], '開発・配置', ['Visual Studio・VS Code', 'dotnet CLI', 'Docker', 'GitHub Actions'])
    legend(s, [(C['app'], 'お客様のアプリ'), (C['ms'], 'Microsoft'), (C['msLight'], 'コミュニティ'), (C['net'], '.NET の実行環境'), (C['core'], 'CoreLib + ランタイム(差し替え不可)'),
               (C['host'], 'プロセスのホスト'), (C['os'], 'OS')])
    return s


def slide3():
    s = Slide('③ Windows / Linux サーバー + FrameworkOnCore')
    r = rows(('APP', 64), ('KINDS', 104), ('CORE', 64), ('LIBS', 152), ('CLASSLIBS', 44), ('CORELIB', 46), ('OS', 50), ('HW', 22))
    app_kinds(s, r, APP_KINDS)
    kinds(s, r, [
        (C['foc'], 'FrameworkOnCore(移植)', ['System.Web.UI(ページ)', 'ASP.NET AJAX(Extensions)', 'Chart・Dynamic Data・Mobile', 'バンドル(Optimization)']),
        (C['foc'], 'MVC 5(移植)', ['System.Web.Mvc', 'Razor・WebPages', 'AspNetWebStack', 'バンドル']),
        (C['foc'], 'Web API 2(移植)', ['Web.Http', 'Http.WebHost', 'Newtonsoft.Json']),
        (C['foc'], 'ASMX(移植)', ['Web.Services', 'WSDL・SOAP']),
        (C['focLight'], 'CoreWCF + FOC', ['CoreWCF', '.svc の構成', 'ASP.NET 互換']),
    ])
    # The application: on System.Web's core, the data libraries only it uses and the ports it shares with System.Web, then
    # on Compat, a partial layer on .NET 10 (the calls whose result differs from .NET Framework's, or that .NET 10 has not),
    # beside a strip down to the class libraries (the others: .NET 10 directly).
    application(s, r, 'お客様のアプリ\n全種類で共通(ほぼそのまま)', ['Global.asax', 'web.config', 'App_Code', 'クラスライブラリ(C#・VB)', '参照する DLL(付け替え)'],
                compat=('FrameworkOnCore\n.Compat', ['数値・カルチャ', 'ファイル・VB', '削除された型']))
    core(s, r, C['foc'], 'FrameworkOnCore の System.Web のコア ― 全種類で共通(referencesource の移植)',
         CORE_ITEMS + ['Compilers(実行時コンパイル + 書き換え)'])
    app_libs(s, r, C['focLight'], 'アプリのデータ', ['EF4 → EF6', 'LINQ to SQL(移植)'])
    # Under System.Web's core, side by side (neither uses the other): the host it runs in (IIS's features as middleware, in
    # IIS's order, on ASP.NET Core and Kestrel), and the ports System.Web and the application use.
    # System.Web's core uses both the middleware (its types: TrySkipIisCustomErrors' flags) and ASP.NET Core (its HttpContext):
    # ASP.NET Core under the middleware and, at its right, up to System.Web's core (one box in two parts, as the application's steps).
    y, middleware_h = r['LIBS'][0], 100
    s.rect(LEFT, y, HOST_W - 54, middleware_h, C['foc'], 'IIS 互換ミドルウェア(FrameworkOnCore。IIS と同じ順、既定は IIS 互換)',
           ['W3C ログ', '応答ヘッダー(customHeaders)', '圧縮', 'エラー ページ(httpErrors)', 'URL Rewrite', 'httpRedirect',
            '既定のドキュメント・静的ファイル・フォルダーの一覧', 'セッション'])
    s.rect(LEFT + HOST_W - 50, y, 50, middleware_h + 4, C['ms'])
    s.rect(LEFT, y + middleware_h + 4, HOST_W, r['LIBS'][1] - middleware_h - 4, C['ms'], 'ASP.NET Core(Microsoft)',
           ['Kestrel(.NET の上のライブラリ)', 'ミドルウェアのパイプライン', '構成・ログ・DI'])
    http_server(s, LEFT + HOST_W - 190, y + middleware_h + 6)
    shared_libs(s, r, C['focLight'], 'ライブラリの移植 ― System.Web もアプリも使う',
                ['SqlClient(移植)', 'Drawing.Common(Linux)', 'BinaryFormatter(互換)', '構成(移植)'], x0=LEFT + HOST_W + 4)
    full(s, r['CLASSLIBS'], C['net'], '.NET 10 クラスライブラリ(そのまま。Compat も、その上で動く)',
         ['System.Data', 'System.Xml', 'System.Text.Json', 'System.Net.Http', 'System.Diagnostics.EventLog'])
    corelib(s, r, 'System.Private.CoreLib + CoreCLR(一組)', ['.NET のまま(ASP.NET Core は .NET の動きが前提)', 'アプリの、結果が違う呼び出しだけ Compat に書き換える'])
    halves(s, r['OS'], (C['os'], 'Windows Server', ['IIS + ANCM(任意)', 'NTFS', 'NLS', 'タイムゾーン']),
           (C['os'], 'Linux', ['Docker・systemd', 'ext4・xfs', 'ICU', 'TZ = 元のサーバーの', 'libgdiplus']))
    full(s, r['HW'], C['hw'], 'サーバー・コンテナ・クラウド(x64・ARM64)', anchor='middle')
    band(s, r['KINDS'][0], bottom(r['LIBS']) - 8, '独自の翻訳(各アセンブリのサテライト)', C['compat'])
    band(s, r['CLASSLIBS'][0] + 8, bottom(r['OS']), 'カルチャ プロファイル', C['compat'])
    host_band(s, r, 'プロセスを載せる(IIS・Docker・systemd)\n作業プロセス(FOC): bin の変更で再起動')
    side(s, C['compat'], 'FrameworkOnCore Studio', ['① 読み込み・元のビルド', '② 解析(使っている API)', '③ 部品の選択', '④ アプリの設定', '⑤ 変換(書き換え規則)',
                                                    '⑥ 配置(Docker・systemd)', 'ParityTest(IIS と比べる)', 'capture-culture・iis'])
    legend(s, [(C['app'], 'お客様のアプリ'), (C['foc'], 'FrameworkOnCore(移植)'), (C['focLight'], '移植・代わり'), (C['ms'], 'Microsoft'),
               (C['compat'], '互換・ツール'), (C['net'], '.NET'), (C['core'], 'CoreLib + ランタイム'), (C['host'], 'ホスト'), (C['os'], 'OS')])
    return s


def slide4():
    s = Slide('CoreLib だけは例外: ランタイムと一組で、アプリから差し替えられない')
    # Left: an ordinary assembly.
    s.rect(14, 46, 590, 330, 'F4F8FC', '', border=C['ms'], border_width=2)
    s.text(26, 52, 560, 26, '普通のアセンブリ(例: System.Web)', size=17, color=C['ms'], bold=True)
    s.rect(60, 92, 200, 50, C['app'], 'アプリ', align='center', anchor='middle', size=15)
    s.rect(330, 92, 230, 50, C['foc'], 'FrameworkOnCore の System.Web\n(referencesource の移植)', align='center', anchor='middle', size=12)
    s.arrow(260, 117, 328, 117, C['compat'])
    s.rect(330, 160, 230, 40, 'FFFFFF', '.NET の System.Web(型転送だけ)', color='999999', size=12, border='BBBBBB', border_width=1.5, dash=True,
           align='center', anchor='middle', bold=False)
    s.text(270, 160, 60, 40, '✕', size=26, color=C['red'], bold=True, align='center')
    s.text(30, 214, 560, 160, '・アプリと一緒に bin に置ける\n・同じ名前でも、アプリが選んだ版が読み込まれる(deps.json)\n'
                               '→ referencesource の ASP.NET を移植して丸ごと差し替えた\n   (意地悪コーパスでも最初から一致)', size=15)
    # Right: CoreLib.
    s.rect(624, 46, 602, 330, 'FBF6EA', '', border=C['gold'], border_width=3)
    s.text(636, 52, 580, 26, 'System.Private.CoreLib(.NET Framework の mscorlib)', size=17, color=C['core'], bold=True)
    s.rect(650, 92, 250, 92, C['core'], 'CoreLib(C#)\nobject・string・double\nCultureInfo・Directory', align='center', anchor='middle', size=13)
    s.rect(900, 92, 250, 92, '2B1A52', 'CoreCLR(C++)\nJIT・GC・型の形', align='center', anchor='middle', size=13)
    s.badge(830, 80, 140, 24, '一組でビルド', fill=C['gold'], color=C['ink'], size=12)
    s.badge(1080, 82, 136, 24, '差し替え不可', size=12)
    s.text(640, 192, 580, 184, '・ランタイムが string・object・Exception の形を直接読む\n・CoreLib がランタイムの C++ の関数を直接呼ぶ\n'
                                '・JIT が一部のメソッドを特別扱いする\n・アプリより先に読み込まれ、1 プロセスに 1 つ(deps.json で選べない)\n'
                                '・mscorlib も CLR 4 と一体(double の書式は clr.dll の C++)', size=15)
    # Bottom: what FrameworkOnCore does.
    s.rect(14, 390, 1212, 266, 'F3FAF5', '', border=C['compat'], border_width=2)
    s.text(26, 396, 1180, 26, 'そこで: CoreLib は .NET のまま、アプリの呼び出しだけを置き換える', size=17, color=C['compat'], bold=True)
    s.rect(40, 440, 230, 40, C['app'], 'アプリのコード', align='center', anchor='middle', size=14)
    s.rect(300, 440, 230, 40, 'FFFFFF', 'price.ToString()', color=C['ink'], border='999999', border_width=1.5, align='center', anchor='middle', size=14)
    s.arrow(532, 460, 600, 460, C['compat'])
    s.text(530, 418, 90, 22, '変換器', size=12, color=C['compat'], bold=True, align='center')
    s.rect(602, 440, 330, 40, C['compat'], 'FrameworkNumbers.ToString(price)', align='center', anchor='middle', size=14)
    s.text(944, 446, 280, 30, '→ 0.3(.NET Framework と同じ)', size=15, bold=True)
    s.rect(40, 500, 230, 40, C['ms'], 'ASP.NET Core・JSON など', align='center', anchor='middle', size=14)
    s.rect(300, 500, 230, 40, 'FFFFFF', 'double.ToString()', color=C['ink'], border='999999', border_width=1.5, align='center', anchor='middle', size=14)
    s.arrow(532, 520, 600, 520, '888888')
    s.text(530, 478, 90, 22, 'そのまま', size=12, color='888888', bold=True, align='center')
    s.rect(602, 500, 330, 40, C['core'], 'CoreLib', align='center', anchor='middle', size=14)
    s.text(944, 506, 280, 30, '→ 0.30000000000000004(.NET の前提どおり)', size=13, bold=True)
    s.text(40, 560, 1170, 90, 'CoreLib を自前で作り直す案は取らない: .NET 全体(毎月のセキュリティ更新、Windows・Linux・x64・ARM64)を自前で保守することになり、\n'
                               'プロセス全体の動きが変わるので .NET 側の部品(ASP.NET Core、System.Text.Json など)を壊すおそれがある。', size=14, color='555555')
    return s


if __name__ == '__main__':
    slides = [slide1(), slide2(), slide3(), slide4()]
    to_pptx(slides, os.path.join(HERE, 'architecture.pptx'))
    for i, s in enumerate(slides, 1):
        to_html(s, os.path.join(os.environ.get('TEMP', HERE), f'architecture-preview-{i}.html'))
    print('ok')
